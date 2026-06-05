"""
Email utilities for sending reports.
"""
import io
from flask import current_app
from flask_mail import Message
from app import mail


def send_report_email(to_addresses: list, subject: str, body: str, attachment_data: bytes = None, attachment_name: str = 'report.xlsx'):
    """Send email with optional Excel attachment."""
    try:
        msg = Message(
            subject=subject,
            recipients=to_addresses,
            body=body,
            sender=current_app.config.get('MAIL_DEFAULT_SENDER', 'noreply@example.com'),
        )
        if attachment_data:
            msg.attach(
                attachment_name,
                'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                attachment_data,
            )
        mail.send(msg)
        return True, 'Email sent successfully.'
    except Exception as e:
        current_app.logger.error(f'Email send error: {e}')
        return False, str(e)


def build_excel_bytes(headers: list, rows: list, sheet_name: str = 'Report') -> bytes:
    """Build an Excel file in memory and return as bytes."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='1a56db')

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].auto_size = True

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
