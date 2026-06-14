"""
Reports blueprint – sales reports, actual vs target, weekly comparison, Excel export.
"""
import io
import calendar
from datetime import date, timedelta

from flask import (
    Blueprint, render_template, request, redirect,
    url_for, flash, send_file, current_app,
)
from flask_babel import lazy_gettext as _l
from sqlalchemy import func, extract, cast, Integer

from app import db
from app.models import (
    Employee, SaleRecord, SalesTarget, CompanyTarget,
    SalesProductivityRecord, AppSetting,
)

bp = Blueprint('reports', __name__)


# ─── Helpers ──────────────────────────────────────────────────────────────────

def _working_days_setting() -> int:
    setting = AppSetting.query.filter_by(key='working_days_per_month').first()
    if setting and setting.value:
        try:
            return int(setting.value)
        except ValueError:
            pass
    return current_app.config.get('WORKING_DAYS_PER_MONTH', 26)


def _employee_month_sales(employee_id: int, year: int, month: int) -> float:
    total = (
        db.session.query(func.sum(SaleRecord.value))
        .filter(
            SaleRecord.employee_id == employee_id,
            extract('year', SaleRecord.sale_date) == year,
            extract('month', SaleRecord.sale_date) == month,
        )
        .scalar()
    )
    return float(total or 0)


def _month_daily_sales(employee_id: int, year: int, month: int) -> list[dict]:
    """Return list of {day, total} for every day of the month that has sales."""
    rows = (
        db.session.query(
            extract('day', SaleRecord.sale_date).label('day'),
            func.sum(SaleRecord.value).label('total'),
        )
        .filter(
            SaleRecord.employee_id == employee_id,
            extract('year', SaleRecord.sale_date) == year,
            extract('month', SaleRecord.sale_date) == month,
        )
        .group_by(extract('day', SaleRecord.sale_date))
        .order_by(extract('day', SaleRecord.sale_date))
        .all()
    )
    return [{'day': int(r.day), 'total': float(r.total or 0)} for r in rows]


def _category_breakdown(employee_id: int, year: int, month: int) -> list[dict]:
    """Return sales totals by product category."""
    rows = (
        db.session.query(
            SaleRecord.product_category,
            func.sum(SaleRecord.value).label('total'),
        )
        .filter(
            SaleRecord.employee_id == employee_id,
            extract('year', SaleRecord.sale_date) == year,
            extract('month', SaleRecord.sale_date) == month,
        )
        .group_by(SaleRecord.product_category)
        .order_by(func.sum(SaleRecord.value).desc())
        .all()
    )
    return [{'category': r.product_category or _l('Uncategorized'), 'total': float(r.total or 0)} for r in rows]


def _weekly_sales(employee_id: int, year: int, month: int) -> list[dict]:
    """Return totals by ISO week number for the given month."""
    first_day = date(year, month, 1)
    last_day = date(year, month, calendar.monthrange(year, month)[1])

    rows = (
        db.session.query(
            SaleRecord.sale_date,
            func.sum(SaleRecord.value).label('total'),
        )
        .filter(
            SaleRecord.employee_id == employee_id,
            SaleRecord.sale_date >= first_day,
            SaleRecord.sale_date <= last_day,
        )
        .group_by(SaleRecord.sale_date)
        .all()
    )

    week_totals: dict[int, float] = {}
    for row in rows:
        week_num = row.sale_date.isocalendar()[1]
        week_totals[week_num] = week_totals.get(week_num, 0.0) + float(row.total or 0)

    return [{'week': w, 'total': t} for w, t in sorted(week_totals.items())]


# ─── Routes ───────────────────────────────────────────────────────────────────

@bp.route('/')
def index():
    """Report selection landing page."""
    today = date.today()
    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
    months = [(i, calendar.month_name[i]) for i in range(1, 13)]
    return render_template(
        'reports/index.html',
        employees=employees,
        months=months,
        today=today,
        page_title=_l('Reports'),
    )


@bp.route('/products')
def products():
    """Product category sales analysis with date and employee filters."""
    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)
    employee_id = request.args.get('employee_id', 0, type=int)
    date_from_str = request.args.get('date_from', '')
    date_to_str = request.args.get('date_to', '')
    mode = request.args.get('mode', 'month')

    if mode == 'range' and date_from_str and date_to_str:
        try:
            d_from = date.fromisoformat(date_from_str)
            d_to = date.fromisoformat(date_to_str)
            if d_from > d_to:
                d_from, d_to = d_to, d_from
        except ValueError:
            mode = 'month'

    if mode != 'range':
        d_from = date(year, month, 1)
        d_to = date(year, month, calendar.monthrange(year, month)[1])

    q = db.session.query(
        SaleRecord.product_category,
        func.sum(SaleRecord.value).label('total'),
        func.sum(SaleRecord.qty).label('qty'),
        func.count(SaleRecord.id).label('cnt'),
        func.count(func.distinct(SaleRecord.employee_id)).label('emp_cnt'),
    ).filter(
        SaleRecord.sale_date >= d_from,
        SaleRecord.sale_date <= d_to,
        SaleRecord.product_category.isnot(None),
        SaleRecord.product_category != '',
    )
    if employee_id:
        q = q.filter(SaleRecord.employee_id == employee_id)

    rows = q.group_by(SaleRecord.product_category).order_by(func.sum(SaleRecord.value).desc()).all()

    grand_total = float(sum(r.total or 0 for r in rows))
    categories = [
        {
            'name': r.product_category,
            'total': float(r.total or 0),
            'qty': float(r.qty or 0),
            'cnt': int(r.cnt),
            'emp_cnt': int(r.emp_cnt),
            'pct': round(float(r.total or 0) / grand_total * 100, 1) if grand_total else 0.0,
        }
        for r in rows
    ]

    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
    months = [(i, calendar.month_name[i]) for i in range(1, 13)]

    return render_template(
        'reports/products.html',
        categories=categories,
        grand_total=grand_total,
        employees=employees,
        employee_id=employee_id,
        d_from=d_from,
        d_to=d_to,
        year=year,
        month=month,
        months=months,
        mode=mode,
        date_from=date_from_str,
        date_to=date_to_str,
        today=today,
        page_title=_l('Product Analysis'),
    )


@bp.route('/employee/<int:emp_id>')
def employee_report(emp_id):
    """Full individual employee sales dashboard."""
    if emp_id == 0:
        first = Employee.query.filter_by(is_active=True).order_by(Employee.name).first()
        if first:
            return redirect(url_for('reports.employee_report', emp_id=first.id,
                                    **request.args.to_dict()))
        return redirect(url_for('reports.index'))
    emp = Employee.query.get_or_404(emp_id)
    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)
    working_days = _working_days_setting()

    target_row = SalesTarget.query.filter_by(
        employee_id=emp.id, year=year, month=month
    ).first()
    target_amount = float(target_row.target_amount) if target_row else 0.0
    emp_working_days = int(target_row.working_days) if target_row else working_days

    achieved = _employee_month_sales(emp.id, year, month)
    remaining = max(target_amount - achieved, 0.0)
    progress_pct = round(achieved / target_amount * 100, 1) if target_amount > 0 else 0.0

    days_elapsed = today.day if (year == today.year and month == today.month) else calendar.monthrange(year, month)[1]
    remaining_days = max(calendar.monthrange(year, month)[1] - days_elapsed, 1)
    daily_target = remaining / remaining_days if remaining_days > 0 else 0.0

    daily_sales = _month_daily_sales(emp.id, year, month)
    categories = _category_breakdown(emp.id, year, month)
    weekly = _weekly_sales(emp.id, year, month)

    # Build cumulative daily chart data
    days_in_month = calendar.monthrange(year, month)[1]
    daily_dict = {d['day']: d['total'] for d in daily_sales}
    chart_labels = list(range(1, days_in_month + 1))
    chart_daily = [daily_dict.get(d, 0.0) for d in chart_labels]
    chart_cumulative = []
    running = 0.0
    for v in chart_daily:
        running += v
        chart_cumulative.append(round(running, 2))

    # Target line for chart
    daily_target_line = round(target_amount / emp_working_days, 2) if emp_working_days > 0 else 0.0
    chart_target_line = [round(daily_target_line * d, 2) for d in chart_labels]

    months = [(i, calendar.month_name[i]) for i in range(1, 13)]
    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()

    return render_template(
        'reports/employee.html',
        emp=emp,
        year=year,
        month=month,
        today=today,
        target_amount=target_amount,
        achieved=achieved,
        remaining=remaining,
        progress_pct=progress_pct,
        daily_target=daily_target,
        remaining_days=remaining_days,
        daily_sales=daily_sales,
        categories=categories,
        weekly=weekly,
        chart_labels=chart_labels,
        chart_daily=chart_daily,
        chart_cumulative=chart_cumulative,
        chart_target_line=chart_target_line,
        months=months,
        employees=employees,
        page_title=_l('Employee Report – %(name)s', name=emp.name),
    )


@bp.route('/comparison')
def comparison():
    """Multi-employee side-by-side comparison."""
    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)

    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
    comparison_data = []

    for emp in employees:
        target_row = SalesTarget.query.filter_by(
            employee_id=emp.id, year=year, month=month
        ).first()
        target_amount = float(target_row.target_amount) if target_row else 0.0
        achieved = _employee_month_sales(emp.id, year, month)
        progress_pct = round(achieved / target_amount * 100, 1) if target_amount > 0 else 0.0
        remaining = max(target_amount - achieved, 0.0)
        if progress_pct >= 100:
            badge = 'success'
        elif progress_pct >= 80:
            badge = 'warning'
        else:
            badge = 'danger'

        comparison_data.append({
            'employee': emp,
            'target_amount': target_amount,
            'achieved': achieved,
            'remaining': remaining,
            'progress_pct': progress_pct,
            'badge': badge,
        })

    comparison_data.sort(key=lambda d: d['achieved'], reverse=True)
    months = [(i, calendar.month_name[i]) for i in range(1, 13)]

    return render_template(
        'reports/comparison.html',
        comparison_data=comparison_data,
        year=year,
        month=month,
        today=today,
        months=months,
        page_title=_l('Employee Comparison'),
    )


@bp.route('/weekly')
def weekly():
    """Weekly performance comparison across all active employees."""
    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)

    first_day = date(year, month, 1)
    last_day = date(year, month, calendar.monthrange(year, month)[1])

    # Get all weeks in the month
    all_weeks: list[int] = []
    d = first_day
    while d <= last_day:
        w = d.isocalendar()[1]
        if w not in all_weeks:
            all_weeks.append(w)
        d += timedelta(days=1)

    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
    weekly_data = []

    for emp in employees:
        week_sales = _weekly_sales(emp.id, year, month)
        week_dict = {w['week']: w['total'] for w in week_sales}
        totals = [week_dict.get(w, 0.0) for w in all_weeks]
        monthly_total = sum(totals)
        weekly_data.append({
            'employee': emp,
            'week_totals': totals,
            'monthly_total': monthly_total,
        })

    weekly_data.sort(key=lambda d: d['monthly_total'], reverse=True)
    months = [(i, calendar.month_name[i]) for i in range(1, 13)]

    return render_template(
        'reports/weekly.html',
        weekly_data=weekly_data,
        all_weeks=all_weeks,
        year=year,
        month=month,
        today=today,
        months=months,
        page_title=_l('Weekly Performance'),
    )


@bp.route('/actual-vs-target')
def actual_vs_target():
    """Achievement report: actual sales vs monthly target per employee."""
    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)
    working_days = _working_days_setting()

    employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
    report_data = []
    grand_target = 0.0
    grand_achieved = 0.0

    for emp in employees:
        target_row = SalesTarget.query.filter_by(
            employee_id=emp.id, year=year, month=month
        ).first()
        target_amount = float(target_row.target_amount) if target_row else 0.0
        achieved = _employee_month_sales(emp.id, year, month)
        variance = achieved - target_amount
        progress_pct = round(achieved / target_amount * 100, 1) if target_amount > 0 else 0.0

        if progress_pct >= 100:
            status = 'achieved'
            badge = 'success'
        elif progress_pct >= 80:
            status = 'near'
            badge = 'warning'
        else:
            status = 'below'
            badge = 'danger'

        grand_target += target_amount
        grand_achieved += achieved

        report_data.append({
            'employee': emp,
            'target_amount': target_amount,
            'achieved': achieved,
            'variance': variance,
            'progress_pct': progress_pct,
            'status': status,
            'badge': badge,
            'percentage': float(target_row.percentage) if target_row else 0.0,
        })

    report_data.sort(key=lambda d: d['progress_pct'], reverse=True)
    grand_pct = round(grand_achieved / grand_target * 100, 1) if grand_target > 0 else 0.0
    months = [(i, calendar.month_name[i]) for i in range(1, 13)]

    return render_template(
        'reports/actual_vs_target.html',
        report_data=report_data,
        grand_target=grand_target,
        grand_achieved=grand_achieved,
        grand_pct=grand_pct,
        year=year,
        month=month,
        today=today,
        months=months,
        page_title=_l('Actual vs Target'),
    )


@bp.route('/export/<report_type>')
def export(report_type):
    """Download Excel report using openpyxl."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, numbers
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash(_l('openpyxl is required for Excel export. Install it with: pip install openpyxl'), 'danger')
        return redirect(url_for('reports.index'))

    today = date.today()
    year = request.args.get('year', today.year, type=int)
    month = request.args.get('month', today.month, type=int)

    wb = openpyxl.Workbook()
    ws = wb.active

    # Common header style
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E79')
    center = Alignment(horizontal='center')

    def _style_header_row(row_idx: int, num_cols: int):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row_idx, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center

    if report_type == 'actual_vs_target':
        ws.title = str(_l('Actual vs Target'))
        headers = [
            str(_l('Employee')), str(_l('Target (SAR)')),
            str(_l('Achieved (SAR)')), str(_l('Variance (SAR)')),
            str(_l('Achievement %')), str(_l('Status')),
        ]
        ws.append(headers)
        _style_header_row(1, len(headers))

        employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
        for emp in employees:
            target_row = SalesTarget.query.filter_by(
                employee_id=emp.id, year=year, month=month
            ).first()
            target_amount = float(target_row.target_amount) if target_row else 0.0
            achieved = _employee_month_sales(emp.id, year, month)
            variance = achieved - target_amount
            pct = round(achieved / target_amount * 100, 1) if target_amount > 0 else 0.0
            status = str(_l('Achieved')) if pct >= 100 else (str(_l('Near')) if pct >= 80 else str(_l('Below')))
            ws.append([emp.name, target_amount, achieved, variance, pct, status])

        filename = f'actual_vs_target_{year}_{month:02d}.xlsx'

    elif report_type == 'comparison':
        ws.title = str(_l('Comparison'))
        headers = [
            str(_l('Employee')), str(_l('Monthly Sales (SAR)')),
            str(_l('Target (SAR)')), str(_l('Achievement %')),
        ]
        ws.append(headers)
        _style_header_row(1, len(headers))

        employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
        for emp in employees:
            target_row = SalesTarget.query.filter_by(
                employee_id=emp.id, year=year, month=month
            ).first()
            target_amount = float(target_row.target_amount) if target_row else 0.0
            achieved = _employee_month_sales(emp.id, year, month)
            pct = round(achieved / target_amount * 100, 1) if target_amount > 0 else 0.0
            ws.append([emp.name, achieved, target_amount, pct])

        filename = f'comparison_{year}_{month:02d}.xlsx'

    elif report_type == 'weekly':
        ws.title = str(_l('Weekly'))
        first_day = date(year, month, 1)
        last_day = date(year, month, calendar.monthrange(year, month)[1])

        all_weeks: list[int] = []
        d = first_day
        while d <= last_day:
            w = d.isocalendar()[1]
            if w not in all_weeks:
                all_weeks.append(w)
            d += timedelta(days=1)

        headers = [str(_l('Employee'))] + [f'W{w}' for w in all_weeks] + [str(_l('Total'))]
        ws.append(headers)
        _style_header_row(1, len(headers))

        employees = Employee.query.filter_by(is_active=True).order_by(Employee.name).all()
        for emp in employees:
            week_sales = _weekly_sales(emp.id, year, month)
            week_dict = {ws_item['week']: ws_item['total'] for ws_item in week_sales}
            row_vals = [emp.name] + [week_dict.get(w, 0.0) for w in all_weeks]
            row_vals.append(sum(row_vals[1:]))
            ws.append(row_vals)

        filename = f'weekly_{year}_{month:02d}.xlsx'

    elif report_type == 'sales_detail':
        ws.title = str(_l('Sales Detail'))
        headers = [
            str(_l('Date')), str(_l('Employee')), str(_l('Invoice No')),
            str(_l('Item Code')), str(_l('Description')), str(_l('Category')),
            str(_l('Qty')), str(_l('Price')), str(_l('Value')),
            str(_l('Ret Qty')), str(_l('Ret Val')),
        ]
        ws.append(headers)
        _style_header_row(1, len(headers))

        records = (
            db.session.query(SaleRecord, Employee.name)
            .join(Employee, Employee.id == SaleRecord.employee_id)
            .filter(
                extract('year', SaleRecord.sale_date) == year,
                extract('month', SaleRecord.sale_date) == month,
            )
            .order_by(SaleRecord.sale_date, Employee.name)
            .all()
        )
        for rec, emp_name in records:
            ws.append([
                rec.sale_date.isoformat() if rec.sale_date else '',
                emp_name,
                rec.invoice_no or '',
                rec.item_code or '',
                rec.description or '',
                rec.product_category or '',
                rec.qty,
                rec.price,
                rec.value,
                rec.ret_qty,
                rec.ret_val,
            ])

        filename = f'sales_detail_{year}_{month:02d}.xlsx'

    else:
        flash(_l('Unknown report type.'), 'danger')
        return redirect(url_for('reports.index'))

    # Auto-fit columns
    for col_cells in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col_cells[0].column)
        for cell in col_cells:
            try:
                cell_len = len(str(cell.value)) if cell.value else 0
                if cell_len > max_len:
                    max_len = cell_len
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    # Write to in-memory buffer
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename,
    )


# ─── Sales Sheet (full SAP-style grouped view + export) ───────────────────────

def _build_sheet_groups(year: int, month: int):
    """Return (emp_order, emp_groups, grand_*) for the sales sheet."""
    first_day = date(year, month, 1)
    last_day  = date(year, month, calendar.monthrange(year, month)[1])

    rows = (
        db.session.query(SaleRecord, Employee.name.label('emp_name'))
        .join(Employee, Employee.id == SaleRecord.employee_id)
        .filter(
            SaleRecord.sale_date >= first_day,
            SaleRecord.sale_date <= last_day,
        )
        .order_by(Employee.name, SaleRecord.sale_date, SaleRecord.invoice_no)
        .all()
    )

    emp_groups: dict[str, dict] = {}
    emp_order:  list[str]       = []

    for rec, emp_name in rows:
        if emp_name not in emp_groups:
            emp_groups[emp_name] = {
                'records':      [],
                'total_gross':  0.0,
                'total_ret':    0.0,
                'total_net':    0.0,
                'total_qty':    0.0,
                'count':        0,
            }
            emp_order.append(emp_name)

        g     = emp_groups[emp_name]
        net   = float(rec.value   or 0)
        ret   = float(rec.ret_val or 0)
        gross = net + ret

        g['records'].append(rec)
        g['total_net']   += net
        g['total_ret']   += ret
        g['total_gross'] += gross
        g['total_qty']   += float(rec.qty or 0)
        g['count']       += 1

    grand_net   = sum(g['total_net']   for g in emp_groups.values())
    grand_ret   = sum(g['total_ret']   for g in emp_groups.values())
    grand_gross = sum(g['total_gross'] for g in emp_groups.values())
    grand_qty   = sum(g['total_qty']   for g in emp_groups.values())
    grand_count = sum(g['count']       for g in emp_groups.values())

    return emp_order, emp_groups, grand_net, grand_ret, grand_gross, grand_qty, grand_count


@bp.route('/sales-sheet')
def sales_sheet():
    """Full SAP-style sales sheet: all invoices grouped by employee with subtotals."""
    today = date.today()
    year  = request.args.get('year',  today.year,  type=int)
    month = request.args.get('month', today.month, type=int)

    emp_order, emp_groups, grand_net, grand_ret, grand_gross, grand_qty, grand_count = \
        _build_sheet_groups(year, month)

    months = [(i, calendar.month_name[i]) for i in range(1, 13)]

    return render_template(
        'reports/sales_sheet.html',
        emp_order=emp_order,
        emp_groups=emp_groups,
        grand_net=grand_net,
        grand_ret=grand_ret,
        grand_gross=grand_gross,
        grand_qty=grand_qty,
        grand_count=grand_count,
        year=year,
        month=month,
        today=today,
        months=months,
        page_title='تقرير الشيت الكامل',
    )


@bp.route('/sales-sheet/export')
def sales_sheet_export():
    """Export the grouped sales sheet to a formatted Excel file."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    today = date.today()
    year  = request.args.get('year',  today.year,  type=int)
    month = request.args.get('month', today.month, type=int)

    emp_order, emp_groups, grand_net, grand_ret, grand_gross, grand_qty, grand_count = \
        _build_sheet_groups(year, month)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'{year}-{month:02d}'
    ws.sheet_view.rightToLeft = True

    # ── Styles ────────────────────────────────────────────────────────────────
    def _font(bold=False, color='000000', size=10, italic=False):
        return Font(bold=bold, color=color, size=size, italic=italic,
                    name='Arial')

    def _fill(hex_color):
        return PatternFill(fill_type='solid', fgColor=hex_color)

    def _border():
        s = Side(style='thin', color='CBD5E1')
        return Border(left=s, right=s, top=s, bottom=s)

    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    right  = Alignment(horizontal='right',  vertical='center')
    left   = Alignment(horizontal='left',   vertical='center')

    COL_HEADERS = [
        '#', 'التاريخ', 'رقم الفاتورة', 'كود البند',
        'الوصف', 'الفئة', 'الكمية', 'السعر',
        'القيمة الإجمالية', 'الخصم / المرتجع', 'الصافي',
        'نوع الفاتورة',
    ]
    NUM_COLS = len(COL_HEADERS)
    COL_WIDTHS = [5, 12, 14, 14, 32, 20, 8, 12, 16, 16, 14, 14]

    # ── Title ─────────────────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=NUM_COLS)
    title_cell = ws.cell(row=1, column=1,
                         value=f'تقرير المبيعات التفصيلي — {year}/{month:02d}')
    title_cell.font      = _font(bold=True, color='FFFFFF', size=13)
    title_cell.fill      = _fill('1E3A5F')
    title_cell.alignment = center
    ws.row_dimensions[1].height = 28

    # ── Column headers ────────────────────────────────────────────────────────
    ws.append(COL_HEADERS)
    for col in range(1, NUM_COLS + 1):
        c = ws.cell(row=2, column=col)
        c.font      = _font(bold=True, color='FFFFFF', size=10)
        c.fill      = _fill('2563EB')
        c.alignment = center
        c.border    = _border()
    ws.row_dimensions[2].height = 22

    row_num = 3
    num_fmt = '#,##0.00'

    for emp_name in emp_order:
        g       = emp_groups[emp_name]
        records = g['records']

        # Employee header
        ws.merge_cells(start_row=row_num, start_column=1,
                       end_row=row_num,   end_column=NUM_COLS)
        emp_cell = ws.cell(
            row=row_num, column=1,
            value=f'  الموظف :  {emp_name}    ({g["count"]} سجل)',
        )
        emp_cell.font      = _font(bold=True, color='FFFFFF', size=11)
        emp_cell.fill      = _fill('1D4ED8')
        emp_cell.alignment = left
        ws.row_dimensions[row_num].height = 20
        row_num += 1

        # Data rows
        for i, rec in enumerate(records, 1):
            net   = float(rec.value   or 0)
            ret   = float(rec.ret_val or 0)
            gross = net + ret
            is_return = net < 0 or ret > 0

            row_data = [
                i,
                rec.sale_date.strftime('%Y-%m-%d') if rec.sale_date else '',
                rec.invoice_no      or '',
                rec.item_code       or '',
                rec.description     or '',
                rec.product_category or '',
                float(rec.qty or 0),
                float(rec.price or 0),
                gross,
                ret,
                net,
                rec.invoice_type    or '',
            ]
            ws.append(row_data)

            row_fill = _fill('FFF8E1') if is_return else _fill('F8FAFC')
            for col in range(1, NUM_COLS + 1):
                c = ws.cell(row=row_num, column=col)
                c.fill   = row_fill
                c.border = _border()
                c.alignment = center if col in (1, 7, 12) else (right if col >= 7 else left)
                if col in (8, 9, 10, 11):
                    c.number_format = num_fmt
            row_num += 1

        # Employee subtotal row
        sub_data = [
            '', f'إجمالي — {emp_name}', '', '', '', '',
            g['total_qty'], '',
            g['total_gross'], g['total_ret'], g['total_net'], '',
        ]
        ws.append(sub_data)
        for col in range(1, NUM_COLS + 1):
            c = ws.cell(row=row_num, column=col)
            c.font      = _font(bold=True, size=10)
            c.fill      = _fill('DBEAFE')
            c.border    = _border()
            c.alignment = right if col >= 7 else left
            if col in (9, 10, 11):
                c.number_format = num_fmt
        ws.row_dimensions[row_num].height = 18
        row_num += 1

        # Empty separator
        ws.append([''] * NUM_COLS)
        row_num += 1

    # Grand total row
    gt_data = [
        '', 'الإجمالي العام', '', '', '', '',
        grand_qty, '',
        grand_gross, grand_ret, grand_net, '',
    ]
    ws.append(gt_data)
    for col in range(1, NUM_COLS + 1):
        c = ws.cell(row=row_num, column=col)
        c.font      = _font(bold=True, color='FFFFFF', size=12)
        c.fill      = _fill('1E3A5F')
        c.border    = _border()
        c.alignment = right if col >= 7 else center
        if col in (9, 10, 11):
            c.number_format = num_fmt
    ws.row_dimensions[row_num].height = 24

    # Column widths
    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze panes below title+header
    ws.freeze_panes = 'A3'

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f'sales_sheet_{year}_{month:02d}.xlsx',
    )

